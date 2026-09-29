"""
Builds a clean, presentation-ready Excel workbook from the sector-level
company CREATION dataset (sv3.build_clean_sector_dataset()'s output) --
one sheet per department, Sector x Year pivot of creation counts, plus
the raw data and a short caveats note.

Usage:
    python scripts/build_sector_openings_workbook.py

Requires sv3.build_clean_sector_dataset() to have already been run at
least once (reads its saved CSV, no API calls here).
"""

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PRIMARY_COLOR = "162B2B"
HEADER_FONT = Font(color="FFFFFF", bold=True)
HEADER_FILL = PatternFill(start_color=PRIMARY_COLOR, end_color=PRIMARY_COLOR, fill_type="solid")

INPUT_CSV = "data/processed/france_creations_by_sector_clean.csv"
OUTPUT_XLSX = "data/manual/company_openings_by_sector.xlsx"

NOTES = [
    ["Company creations by sector - Bas-Rhin / Haut-Rhin / Moselle"],
    [],
    ["Source: INSEE Sirene v3.11, enterprise-level (etablissementSiege:true + dateCreationUniteLegale),"],
    ["not raw etablissement creations -- see sirene_v3_client.py for why that filter matters."],
    [],
    ["Excludes non-business legal forms: VAT-only registrations, public bodies, indivisions,"],
    ["professional orders, and cooperative unions (5459/5559/5659/6318) -- see"],
    ["data/manual/legal_form_category_review.csv for the full exclusion list."],
    [],
    ["A small share of records ('Unknown / unclassified' sector) carry no NAF code yet at fetch"],
    ["time -- mostly very recent registrations still pending classification by INSEE, not a data"],
    ["quality defect. Check that share before treating the sector breakdown as complete."],
]


def build():
    df = pd.read_csv(INPUT_CSV)

    unclassified_share = (
        df.loc[df["sector"] == "Unknown / unclassified", "count"].sum() / df["count"].sum() * 100
    )

    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        pd.DataFrame(NOTES + [[], [f"Unknown/unclassified share of total: {unclassified_share:.2f}%"]]).to_excel(
            writer, sheet_name="Notes", index=False, header=False
        )

        for region in sorted(df["region"].unique()):
            region_df = df[df["region"] == region]
            pivot = (
                region_df.pivot_table(index="sector", columns="year", values="count", aggfunc="sum", fill_value=0)
                .sort_index()
            )
            sheet_name = region[:31]
            pivot.to_excel(writer, sheet_name=sheet_name)

        df.to_excel(writer, sheet_name="Raw data", index=False)

    wb = load_workbook(OUTPUT_XLSX)
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        if sheet_name == "Notes":
            ws.column_dimensions["A"].width = 100
            continue

        for cell in ws[1]:
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal="center")

        ws.column_dimensions["A"].width = 35
        ws.freeze_panes = "B2"
        for col_idx in range(2, ws.max_column + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 12

    wb.save(OUTPUT_XLSX)
    print(f"Saved: {OUTPUT_XLSX}")
    print(f"Unknown/unclassified share of total: {unclassified_share:.2f}%")


if __name__ == "__main__":
    build()
